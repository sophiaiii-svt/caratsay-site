import openpyxl, sys
f = r"C:\Users\EDY\Documents\xwechat_files\wxid_d2tswczcilzq12_ccac\temp\RWTemp\2026-09\2f6b875213ac3f03c98bae6fec0cddfc\Going_Seventeen_节目列表.xlsx"
wb = openpyxl.load_workbook(f, data_only=True)
for ws in wb.worksheets:
    print(f"SHEET: {ws.title}  rows={ws.max_row} cols={ws.max_column}")
    for r, row in enumerate(ws.iter_rows(values_only=True), 1):
        cells = [("" if c is None else str(c)).strip() for c in row]
        while cells and cells[-1] == "":
            cells.pop()
        if not any(cells):
            continue
        print(f"[{r:03d}] " + " | ".join(cells))
